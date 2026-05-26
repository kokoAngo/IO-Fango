#!/usr/bin/env python3
"""Render listing PDFs into thumbnails + extract docx field corpus.

Inputs: data/raw/*.pdf and data/raw/*.docx (optional; no-op if absent).
Outputs:
  * data/thumbnails/<stem>.png — first-page render
  * data/seed/listing-rain.json — flat list of field strings extracted from any docx

Soft dependencies: pypdfium2 (PDF), python-docx (DOCX). If missing, that part is
skipped with a warning.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from fango.config import DATA_DIR, SEED_DIR

RAW_DIR = DATA_DIR / "raw"
THUMB_DIR = DATA_DIR / "thumbnails"


def render_pdfs() -> int:
    if not RAW_DIR.exists():
        return 0
    try:
        import pypdfium2 as pdfium  # type: ignore
    except ImportError:
        print("pypdfium2 not installed; skipping PDF rendering", file=sys.stderr)
        return 0
    THUMB_DIR.mkdir(parents=True, exist_ok=True)
    n = 0
    for pdf in sorted(RAW_DIR.glob("*.pdf")):
        try:
            doc = pdfium.PdfDocument(str(pdf))
            page = doc[0]
            pil = page.render(scale=1.5).to_pil()
            out = THUMB_DIR / f"{pdf.stem}.png"
            pil.save(out)
            n += 1
            print(f"rendered {pdf.name} -> {out}")
        except Exception as exc:
            print(f"failed {pdf}: {exc}", file=sys.stderr)
    return n


def extract_docx_corpus() -> int:
    if not RAW_DIR.exists():
        SEED_DIR.mkdir(parents=True, exist_ok=True)
        (SEED_DIR / "listing-rain.json").write_text("[]", encoding="utf-8")
        return 0
    try:
        import docx  # type: ignore
    except ImportError:
        print("python-docx not installed; skipping DOCX extraction", file=sys.stderr)
        SEED_DIR.mkdir(parents=True, exist_ok=True)
        (SEED_DIR / "listing-rain.json").write_text("[]", encoding="utf-8")
        return 0
    fields: list[str] = []
    for path in sorted(RAW_DIR.glob("*.docx")):
        try:
            d = docx.Document(str(path))
            for para in d.paragraphs:
                t = para.text.strip()
                if t:
                    fields.append(t)
        except Exception as exc:
            print(f"failed {path}: {exc}", file=sys.stderr)
    SEED_DIR.mkdir(parents=True, exist_ok=True)
    out = SEED_DIR / "listing-rain.json"
    out.write_text(json.dumps(fields, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"extracted {len(fields)} field strings -> {out}")
    return len(fields)


def main() -> int:
    n_pdf = render_pdfs()
    n_fields = extract_docx_corpus()
    print(f"done: {n_pdf} PDFs rendered, {n_fields} docx field strings")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
